"""Génère le support PowerPoint (16:9) avec les notes d'orateur."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import content as C
from slides import ANNEXES, BUDGET, CRITERES, SLIDES

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "img"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
ORANGE = RGBColor(0xE7, 0x6F, 0x51)
GOLD = RGBColor(0xE9, 0xC4, 0x6A)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x11, 0x18, 0x27)

SW, SH = Inches(13.333), Inches(7.5)

MOSCOW_COURT = {
    "Must have": ["✓ Données point-in-time, licences vérifiées",
                  "✓ 4 régimes sur 3 dimensions + jump model",
                  "✓ Détecteur de stress v2 pré-enregistré",
                  "✓ Probabilité de stress calibrée",
                  "✓ Explication du signal",
                  "✓ Note de comité PDF, FR / EN",
                  "✓ Track record public horodaté",
                  "☐ Tests avec ≥ 3 gérants (J3)"],
    "Should have": ["◐ Ralentissement fiable (0,492 auj.)",
                    "☐ Overlay de risque 60/40",
                    "☐ Couche actualité en challenger",
                    "✓ Alertes et scénarios de stress",
                    "✓ API REST"],
    "Could have": ["✓ Zones UK, Japon, émergents (euro échouée)",
                   "✓ Pages « Apprendre », carte mondiale",
                   "☐ Tuteur IA GAMA",
                   "✓ Rapport de thèse généré"],
    "Won't have": ["Signaux achat / vente (MiFID II)",
                   "Vente / redistribution de données",
                   "Deep learning, SOM, TDA (mémoire)",
                   "Temps réel / intraday",
                   "Comptes, paiement, app mobile"],
}
MOSCOW_COL = {"Must have": RED, "Should have": ORANGE, "Could have": TEAL, "Won't have": GREY}


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=16, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullet=False, spacing=6, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = ("• " if bullet else "") + ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return tb


def picture(slide, name, x, y, max_w, max_h):
    path = IMG / name
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    return slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y + (max_h - h) // 2, w, h)


def frame(prs, title, num, orateur=None, duree=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, Inches(0.95), NAVY)
    rect(s, 0, Inches(0.95), SW, Inches(0.06), ORANGE)
    text(s, Inches(0.45), Inches(0.12), Inches(11.4), Inches(0.75), title, size=26,
         color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(0.45), Inches(7.08), Inches(9), Inches(0.35),
         f"{C.GROUPE} · {C.PROJET}", size=10, color=GREY)
    text(s, Inches(12.2), Inches(7.08), Inches(0.9), Inches(0.35), str(num), size=11,
         color=GREY, align=PP_ALIGN.RIGHT)
    if orateur:
        tag = rect(s, Inches(11.55), Inches(0.28), Inches(1.6), Inches(0.42), TEAL,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = tag.text_frame
        tf.text = f"{orateur} · {duree}s"
        tf.paragraphs[0].runs[0].font.size = Pt(11)
        tf.paragraphs[0].runs[0].font.color.rgb = WHITE
        tf.paragraphs[0].runs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return s


def notes(slide, sd):
    tf = slide.notes_slide.notes_text_frame
    tf.text = (f"ORATEUR : {sd['orateur']} — durée cible : {sd['duree']} s\n\n"
               f"{sd['script']}")


def bullets_panel(s, items, x, y, w, h, size=15):
    rect(s, x, y, w, h, LIGHT)
    rect(s, x, y, Inches(0.08), h, ORANGE)
    text(s, x + Inches(0.25), y + Inches(0.2), w - Inches(0.4), h - Inches(0.3), items,
         size=size, bullet=True, spacing=10)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    total = 0
    for n, sd in enumerate(SLIDES, 1):
        total += sd["duree"]
        k = sd["kind"]
        if k == "title":
            s = prs.slides.add_slide(prs.slide_layouts[6])
            rect(s, 0, 0, SW, SH, NAVY)
            rect(s, 0, Inches(4.55), SW, Inches(0.08), ORANGE)
            text(s, Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.4), sd["titre"], size=44,
                 color=WHITE, bold=True)
            text(s, Inches(0.8), Inches(2.75), Inches(11.5), Inches(1.4),
                 "Aider les petits gérants européens à réduire le risque plus tôt "
                 "— et à le justifier", size=22, color=GOLD)
            text(s, Inches(0.8), Inches(3.75), Inches(11.5), Inches(0.6),
                 "Soutenance intermédiaire — Jalon 2 · Conception de la solution", size=16,
                 color=WHITE)
            text(s, Inches(0.8), Inches(4.9), Inches(11.5), Inches(1.4),
                 [f"{C.GROUPE}",
                  " · ".join(m[0] for m in C.EQUIPE),
                  f"Coach : {C.COACH} · ECE Paris ING5 · 2026-2027"], size=15, color=WHITE)
        elif k == "image":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            picture(s, sd["img"], Inches(0.35), Inches(1.25), Inches(8.6), Inches(5.6))
            bullets_panel(s, sd["bullets"], Inches(9.15), Inches(1.45), Inches(3.85), Inches(5.2))
        elif k == "image_full":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            picture(s, sd["img"], Inches(0.3), Inches(1.15), Inches(12.7), Inches(5.85))
        elif k == "statement":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            text(s, Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.5), sd["statement"],
                 size=22, color=NAVY, bold=True)
            rect(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.5), LIGHT)
            text(s, Inches(1.15), Inches(3.1), Inches(11.0), Inches(0.4), "QUESTION DIRECTRICE",
                 size=12, color=ORANGE, bold=True)
            text(s, Inches(1.15), Inches(3.45), Inches(11.0), Inches(1.0), sd["question"],
                 size=18, color=BLACK)
            for i, (v, lab, res) in enumerate(sd["kpis"]):
                x = Inches(0.9 + i * 3.95)
                rect(s, x, Inches(4.75), Inches(3.6), Inches(1.45), NAVY,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(s, x, Inches(4.8), Inches(3.6), Inches(0.8), v, size=34, color=GOLD,
                     bold=True, align=PP_ALIGN.CENTER)
                text(s, x, Inches(5.6), Inches(3.6), Inches(0.5), lab, size=14, color=WHITE,
                     align=PP_ALIGN.CENTER)
                rect(s, x, Inches(6.25), Inches(3.6), Inches(0.6), RGBColor(0xD1, 0xFA, 0xE5),
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(s, x, Inches(6.25), Inches(3.6), Inches(0.6), res, size=14, color=GREEN,
                     bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif k == "persona":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            rect(s, Inches(0.4), Inches(1.35), Inches(4.3), Inches(5.4), LIGHT)
            ic = rect(s, Inches(0.7), Inches(1.6), Inches(0.9), Inches(0.9), TEAL, shape=MSO_SHAPE.OVAL)
            ic.text_frame.text = "C"
            ic.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
            ic.text_frame.paragraphs[0].runs[0].font.bold = True
            ic.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            text(s, Inches(1.75), Inches(1.7), Inches(2.9), Inches(0.8), "Persona (fictif)",
                 size=12, color=GREY)
            text(s, Inches(1.75), Inches(1.95), Inches(2.9), Inches(0.8), sd["persona"][0],
                 size=16, color=NAVY, bold=True)
            text(s, Inches(0.65), Inches(2.75), Inches(3.9), Inches(3.9), sd["persona"][1:],
                 size=14, bullet=True, spacing=9)
            picture(s, sd["img"], Inches(4.95), Inches(1.3), Inches(8.1), Inches(4.3))
            for i, b in enumerate(sd["bullets"]):
                x = Inches(5.1 + i * 4.0)
                rect(s, x, Inches(5.85), Inches(3.8), Inches(0.85), NAVY if i == 0 else ORANGE,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(s, x, Inches(5.85), Inches(3.8), Inches(0.85), b, size=15, color=WHITE,
                     bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif k == "solution":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            text(s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(1.0), C.PITCH_2_PHRASES,
                 size=15, color=NAVY, bold=True)
            pic = picture(s, sd["img"], Inches(0.35), Inches(2.3), Inches(6.9), Inches(4.45))
            pic.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
            text(s, Inches(0.35), Inches(6.78), Inches(6.9), Inches(0.3),
                 "Page publique réelle du track record (track_record/fr.html, 20/10/2026)",
                 size=9, color=GREY, align=PP_ALIGN.CENTER)
            text(s, Inches(7.5), Inches(2.3), Inches(5.5), Inches(0.4), "ÉVOLUTIONS DEPUIS J1",
                 size=12, color=TEAL, bold=True)
            text(s, Inches(7.5), Inches(2.65), Inches(5.5), Inches(2.3),
                 ["3 dimensions (stress, croissance, inflation) au lieu de 2",
                  "Détecteur de stress v2 pré-enregistré : 11/11 crises",
                  "Suivi live dès le 01/10 (prévu le 30/10), 4 zones",
                  "Couche actualité IA : reportée en Phase 3"],
                 size=12.5, bullet=True, spacing=5)
            rect(s, Inches(7.5), Inches(4.85), Inches(5.5), Inches(2.1), RGBColor(0xFE, 0xF3, 0xC7),
                 line=ORANGE)
            text(s, Inches(7.65), Inches(4.9), Inches(5.2), Inches(0.4),
                 "CE QUI RESTE LE PLUS RISQUÉ", size=12, color=ORANGE, bold=True)
            text(s, Inches(7.65), Inches(5.25), Inches(5.2), Inches(1.7), C.HYPOTHESE_RISQUEE,
                 size=12.5)
        elif k == "moscow":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            colw = Inches(3.08)
            for i, (cat, items) in enumerate(MOSCOW_COURT.items()):
                x = Inches(0.35) + i * (colw + Inches(0.1))
                rect(s, x, Inches(1.25), colw, Inches(0.6), MOSCOW_COL[cat])
                text(s, x, Inches(1.25), colw, Inches(0.6), f"{cat} ({len(items)})", size=18,
                     color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                rect(s, x, Inches(1.85), colw, Inches(4.55), LIGHT)
                text(s, x + Inches(0.1), Inches(1.97), colw - Inches(0.2), Inches(4.4), items,
                     size=13, spacing=8)
            text(s, Inches(0.35), Inches(6.45), Inches(12.6), Inches(0.6),
                 "✓ fait   ◐ en cours   ☐ à faire  (état au 21/10)  —  Écart assumé : des Could ont été "
                 "construits avant des Should → gel des nouvelles fonctionnalités jusqu'aux tests.",
                 size=12.5, color=NAVY, bold=True)
        elif k == "risques":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            picture(s, sd["img"], Inches(0.3), Inches(1.2), Inches(4.6), Inches(5.7))
            majeurs = {r[0]: r for r in C.RISQUES}
            court = [
                ("R1", "Aucun test utilisateur fait : gérants indisponibles en novembre",
                 "Sessions de 30 min planifiées dès le 09/11 sur la note de comité.",
                 "Plan B : protocole asynchrone (vidéo + formulaire de 10 min), CGP, alumni ECE.",
                 "Déclencheur : < 3 sessions confirmées au 06/11"),
                ("R2", "En réel, plus de 1,5 fausse alerte par an (1,26 en backtest)",
                 "Règle pré-enregistrée, jamais retouchée ; challengers publiés en live.",
                 "Plan B écrit : détecteur v1 (0,17 fausse alerte/an, 11 j de latence).",
                 "Déclencheur : 2 fausses alertes dans le track record live"),
                ("R3", "Code écrit en grande partie avec un assistant IA : les 5 doivent le défendre",
                 "Un module « propriétaire » par membre, présenté aux autres en 15 min.",
                 "Revue de code obligatoire, répétition des questions techniques.",
                 "Déclencheur : un membre ne sait pas expliquer son module"),
            ]
            top = Inches(1.25)
            for rid, titre, mitig, planb, decl in court:
                r = majeurs[rid]
                rect(s, Inches(5.1), top, Inches(7.9), Inches(1.8), LIGHT)
                rect(s, Inches(5.1), top, Inches(0.9), Inches(1.8), NAVY)
                text(s, Inches(5.1), top, Inches(0.9), Inches(1.8), rid, size=22, color=WHITE,
                     bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(s, Inches(6.1), top + Inches(0.05), Inches(6.8), Inches(0.4), titre,
                     size=13, color=NAVY, bold=True)
                text(s, Inches(6.1), top + Inches(0.42), Inches(6.8), Inches(0.3),
                     f"{r[1]} · Probabilité {r[3].lower()} · Impact {r[4].lower()}",
                     size=10.5, color=GREY)
                text(s, Inches(6.1), top + Inches(0.75), Inches(6.8), Inches(1.0),
                     [mitig, planb, decl], size=11.5, spacing=2)
                top += Inches(1.92)
        elif k == "fin":
            s = frame(prs, sd["titre"], n, sd["orateur"], sd["duree"])
            picture(s, sd["img"], Inches(0.3), Inches(1.15), Inches(8.4), Inches(4.0))
            rect(s, Inches(8.9), Inches(1.3), Inches(4.1), Inches(3.7), NAVY)
            text(s, Inches(9.1), Inches(1.4), Inches(3.8), Inches(0.5), "EN RÉSUMÉ", size=13,
                 color=GOLD, bold=True)
            text(s, Inches(9.1), Inches(1.85), Inches(3.8), Inches(3.1),
                 ["Besoin confirmé par 13 entretiens + 54 réponses",
                  "Choix tracés par 4 arbres et des protocoles pré-enregistrés",
                  "7 risques, chacun avec un plan B et un déclencheur"],
                 size=14, color=WHITE, bullet=True, spacing=12)
            rect(s, Inches(0.4), Inches(5.35), Inches(12.6), Inches(1.55), RGBColor(0xFE, 0xF3, 0xC7),
                 line=ORANGE)
            text(s, Inches(0.6), Inches(5.42), Inches(12.2), Inches(0.4),
                 "VOS RETOURS NOUS INTÉRESSENT SUR :", size=13, color=ORANGE, bold=True)
            text(s, Inches(0.6), Inches(5.8), Inches(12.2), Inches(1.1),
                 ["Une fausse alerte par an est-elle acceptable pour un gérant ?",
                  "Notre protocole de test utilisateur (30 min sur la note de comité) est-il assez exigeant ?"],
                 size=14, bullet=True)
        notes(s, sd)

    # ---------- slide de transition + annexes
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, NAVY)
    text(s, Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.2), "Merci — vos questions",
         size=44, color=WHITE, bold=True)
    text(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.8),
         "Annexes : latence par épisode, arbres n°3 et n°4, arbre actualité, marché, prix, "
         "risques détaillés, critères de succès, budget, état du code", size=16, color=GOLD)
    s.notes_slide.notes_text_frame.text = (
        "Rester sur cette slide pendant les questions. Aller chercher l'annexe utile "
        "seulement si la question le demande (taper le numéro de la slide puis Entrée "
        "en mode diaporama).")
    n = len(SLIDES) + 1
    for a in ANNEXES:
        n += 1
        s = frame(prs, a["titre"], n)
        kind = a.get("kind")
        if a.get("img"):
            picture(s, a["img"], Inches(0.3), Inches(1.15), Inches(12.7), Inches(5.85))
        elif kind == "risk_table":
            table(s, ["ID", "Catégorie", "Risque", "P", "I", "Plan de mitigation", "Déclencheur"],
                  [list(r[:5]) + [r[5], r[6]] for r in C.RISQUES],
                  [0.5, 1.2, 2.6, 0.8, 0.6, 4.9, 2.1], size=9)
        elif kind == "kpi_table":
            table(s, ["Critère", "Objectif (J1)", "Résultat au 21/10"], CRITERES,
                  [4.6, 3.4, 4.6], size=12)
        elif kind == "budget":
            table(s, ["Poste", "Détail", "Coût"], BUDGET, [3.8, 6.5, 2.3], size=14)
        elif kind == "repo":
            text(s, Inches(0.6), Inches(1.4), Inches(6.3), Inches(5.4), [
                "Application Streamlit 7 écrans, API REST FastAPI, ligne de commande",
                "5 modèles : k-means, jump model, gradient boosting, onset, combined",
                "Validation : walk-forward, holdout, calibration, pré-enregistrements (docs/)",
                "Publication quotidienne : 14 jours US depuis le 01/10 + UK, Japon, émergents",
                "29 fichiers de tests (pytest) + ruff, intégration continue",
                "La soutenance intermédiaire n'exige pas de prototype : à montrer seulement si on le demande",
            ], size=15, bullet=True, spacing=12)
            rect(s, Inches(7.2), Inches(1.4), Inches(5.7), Inches(5.3), RGBColor(0x11, 0x18, 0x27))
            text(s, Inches(7.4), Inches(1.55), Inches(5.4), Inches(5.0), [
                "PFE_DRAI/",
                "├── app/            # Streamlit",
                "├── api/            # FastAPI",
                "├── config/         # settings, zones",
                "├── docs/           # protocoles, thèse",
                "├── pfe_drai/",
                "│   ├── data/       # FRED, Tiingo…",
                "│   ├── features/   # 3 scores",
                "│   ├── models/     # jump, onset…",
                "│   ├── validation/ # walk-forward",
                "│   └── publish/    # track record",
                "├── track_record/   # JSON + .ots",
                "└── tests/          # pytest",
            ], size=13, color=RGBColor(0xA7, 0xF3, 0xD0), spacing=1, font="Consolas")
        s.notes_slide.notes_text_frame.text = f"À montrer seulement si : {a['quand']}"

    out = ROOT / "Support_soutenance_blanche_J2.pptx"
    prs.save(out)
    print("pptx ok", out, "durée totale", total, "s")


def table(slide, header, rows, widths, size=11):
    shp = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.35), Inches(1.25),
                                 Inches(sum(widths)), Inches(0.5) * (len(rows) + 1))
    t = shp.table
    for j, w in enumerate(widths):
        t.columns[j].width = Emu(int(Inches(w)))
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        para = c.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(size + 1)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = WHITE
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            for para in c.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(size)
                    r.font.color.rgb = BLACK
                    if i == len(rows) and "Total" in str(row[0]):
                        r.font.bold = True


if __name__ == "__main__":
    build()
